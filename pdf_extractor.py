import os
import re

import fitz
import pytesseract

from PIL import Image
from docx import Document


# ---------------------------------------------------------
# TEXT CLEANING
# ---------------------------------------------------------

def clean_text(text: str) -> str:
    """
    Clean extracted/OCR text while preserving paragraphs.
    """

    if not text:
        return ""

    text = text.replace("\x00", "")
    text = text.replace("“", '"')
    text = text.replace("”", '"')
    text = text.replace("’", "'")

    # Normalize excessive spaces but preserve newlines
    text = re.sub(r"[ \t]{2,}", " ", text)

    # Avoid huge numbers of blank lines
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


# ---------------------------------------------------------
# PDF TEXT EXTRACTION
# ---------------------------------------------------------

def extract_text_from_pdf(pdf_path: str) -> str:
    """
    Extract text from a PDF.

    First tries normal PDF text extraction.
    If the PDF contains little/no usable text,
    automatically falls back to OCR.
    """

    try:
        pdf = fitz.open(pdf_path)

        pages_text = []

        for page in pdf:
            text = page.get_text("text")

            if text:
                pages_text.append(text)

        pdf.close()

        extracted_text = clean_text("\n".join(pages_text))

        # Check whether extracted text is actually useful
        if is_text_usable(extracted_text):
            return extracted_text

        # Otherwise use OCR
        return ocr_pdf(pdf_path)

    except Exception as e:
        raise ValueError(
            f"Could not read PDF: {str(e)}"
        )


# ---------------------------------------------------------
# SCANNED PDF OCR
# ---------------------------------------------------------

def ocr_pdf(pdf_path: str) -> str:
    """
    Convert PDF pages into images and run OCR.
    Supports English + Marathi.
    """

    try:
        pdf = fitz.open(pdf_path)

        page_texts = []

        for page_number, page in enumerate(pdf):

            # Render page at good OCR resolution
            matrix = fitz.Matrix(2.0, 2.0)

            pix = page.get_pixmap(
                matrix=matrix,
                alpha=False
            )

            image = Image.frombytes(
                "RGB",
                [pix.width, pix.height],
                pix.samples
            )

            text = pytesseract.image_to_string(
                image,
                lang="eng+mar"
            )

            if text.strip():
                page_texts.append(
                    f"[Page {page_number + 1}]\n{text}"
                )

        pdf.close()

        result = clean_text(
            "\n\n".join(page_texts)
        )

        if not result:
            raise ValueError(
                "OCR could not extract readable text from the PDF."
            )

        return result

    except Exception as e:
        raise ValueError(
            f"Could not OCR PDF: {str(e)}"
        )


# ---------------------------------------------------------
# DOCX EXTRACTION
# ---------------------------------------------------------

def extract_text_from_docx(docx_path: str) -> str:
    """
    Extract text from a DOCX document.
    """

    try:
        document = Document(docx_path)

        paragraphs = []

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        # Also extract table content
        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    cell_text = cell.text.strip()

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    paragraphs.append(
                        " | ".join(cells)
                    )

        result = clean_text(
            "\n\n".join(paragraphs)
        )

        if not result:
            raise ValueError(
                "No readable text found in DOCX."
            )

        return result

    except Exception as e:
        raise ValueError(
            f"Could not read DOCX: {str(e)}"
        )


# ---------------------------------------------------------
# IMAGE OCR
# ---------------------------------------------------------

def extract_text_from_image(image_path: str) -> str:
    """
    Extract text from JPG/PNG images using Tesseract.
    """

    try:

        image = Image.open(image_path)

        # Convert to RGB for consistent OCR
        image = image.convert("RGB")

        text = pytesseract.image_to_string(
            image,
            lang="eng+mar"
        )

        result = clean_text(text)

        if not result:
            raise ValueError(
                "OCR could not extract readable text from the image."
            )

        return result

    except Exception as e:
        raise ValueError(
            f"Could not process image: {str(e)}"
        )


# ---------------------------------------------------------
# TEXT QUALITY CHECK
# ---------------------------------------------------------

def is_text_usable(text: str) -> bool:
    """
    Determine whether extracted PDF text is usable.

    This prevents scanned PDFs with a few random characters
    from bypassing OCR.
    """

    if not text:
        return False

    text = text.strip()

    if len(text) < 30:
        return False

    # Count letters from English/Marathi/Unicode scripts
    letters = sum(
        1 for char in text
        if char.isalpha()
    )

    if letters < 20:
        return False

    # Require a reasonable amount of readable content
    ratio = letters / max(len(text), 1)

    return ratio >= 0.20


# ---------------------------------------------------------
# UNIVERSAL DOCUMENT EXTRACTION
# ---------------------------------------------------------

def extract_text_from_document(file_path: str) -> str:
    """
    Universal document extractor.

    Supported:
        PDF
        DOCX
        JPG
        JPEG
        PNG
    """

    extension = os.path.splitext(
        file_path
    )[1].lower()

    if extension == ".pdf":

        return extract_text_from_pdf(
            file_path
        )

    elif extension == ".docx":

        return extract_text_from_docx(
            file_path
        )

    elif extension in [
        ".jpg",
        ".jpeg",
        ".png"
    ]:

        return extract_text_from_image(
            file_path
        )

    else:

        raise ValueError(
            "Unsupported file type. "
            "Please upload a PDF, DOCX, JPG, JPEG, or PNG file."
        )


# ---------------------------------------------------------
# BACKWARD COMPATIBILITY
# ---------------------------------------------------------

# Existing code can continue using:
#
# extract_text_from_pdf(...)
#
# while the new backend can use:
#
# extract_text_from_document(...)


if __name__ == "__main__":

    print(
        "LegalEase document extractor loaded successfully."
    )
